#!/usr/bin/env python3
"""Generate a Word (.docx) mock-up of the subdivision analysis report."""

import json
import os
import sys
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor, Cm

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(REPO_ROOT, "_data", "realtyrates")
OUT_PATH = os.path.join(REPO_ROOT, "report", "subdivision-analysis.docx")

# ── Colour palette ──────────────────────────────────────────────────────────
DARK_BLUE  = RGBColor(0x1A, 0x2A, 0x4A)   # headings
MID_BLUE   = RGBColor(0x2E, 0x6D, 0xBF)   # accents / table headers
LIGHT_BLUE = RGBColor(0xD6, 0xE4, 0xF7)   # table header fill
ACCENT_GRN = RGBColor(0x1E, 0x7E, 0x52)   # positive / avg values
ACCENT_ORG = RGBColor(0xC0, 0x50, 0x20)   # interest rate colour
TEXT_GREY  = RGBColor(0x44, 0x44, 0x44)
LIGHT_GREY = RGBColor(0xF2, 0xF4, 0xF8)   # alternating row fill
MID_GREY   = RGBColor(0xCC, 0xCC, 0xCC)   # borders
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)


# ── XML helpers ──────────────────────────────────────────────────────────────

def hex_color(rgb: RGBColor) -> str:
    return f"{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}"


def set_cell_bg(cell, rgb: RGBColor):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color(rgb))
    tcPr.append(shd)


def set_cell_border(cell, **kwargs):
    """kwargs: top, bottom, left, right — each a dict with 'val', 'sz', 'color'."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for side, attrs in kwargs.items():
        border = OxmlElement(f"w:{side}")
        border.set(qn("w:val"),   attrs.get("val",   "single"))
        border.set(qn("w:sz"),    str(attrs.get("sz",  4)))
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), attrs.get("color", "CCCCCC"))
        tcBorders.append(border)
    tcPr.append(tcBorders)


def shade_row(row, rgb: RGBColor):
    for cell in row.cells:
        set_cell_bg(cell, rgb)


def fmt_pct(val, decimals=2):
    if val is None:
        return "—"
    return f"{val:.{decimals}f}%"


def fmt_range(lo, hi, unit="%"):
    if lo is None and hi is None:
        return "—"
    if lo == hi or hi is None:
        return f"{lo:.2f}{unit}"
    return f"{lo:.2f} – {hi:.2f}{unit}"


# ── Document helpers ─────────────────────────────────────────────────────────

def set_font(run, bold=False, italic=False, size=None, color=None):
    run.bold = bold
    run.italic = italic
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color


def heading(doc, text, level=1, color=DARK_BLUE, size=None, after=6):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_after = Pt(after)
    for run in p.runs:
        run.font.color.rgb = color
        if size:
            run.font.size = Pt(size)
    return p


def body(doc, text, size=10, color=TEXT_GREY, italic=False, after=4, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    run = p.add_run(text)
    set_font(run, bold=bold, italic=italic, size=size, color=color)
    return p


def caption(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    set_font(run, italic=True, size=8.5, color=RGBColor(0x77, 0x77, 0x77))
    return p


def hr(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "2E6DBF")
    pBdr.append(bottom)
    pPr.append(pBdr)


def stat_table(doc, stats):
    """Render a 4-column summary stat block."""
    cols = min(len(stats), 4)
    tbl = doc.add_table(rows=3, cols=cols)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, (label, value, sub) in enumerate(stats[:cols]):
        # header row
        hcell = tbl.rows[0].cells[i]
        set_cell_bg(hcell, LIGHT_BLUE)
        p = hcell.paragraphs[0]
        run = p.add_run(label)
        set_font(run, bold=True, size=8, color=MID_BLUE)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # value row
        vcell = tbl.rows[1].cells[i]
        set_cell_bg(vcell, WHITE)
        vp = vcell.paragraphs[0]
        vrun = vp.add_run(value)
        set_font(vrun, bold=True, size=16, color=DARK_BLUE)
        vp.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # sub row
        scell = tbl.rows[2].cells[i]
        set_cell_bg(scell, LIGHT_GREY)
        sp = scell.paragraphs[0]
        srun = sp.add_run(sub)
        set_font(srun, italic=True, size=8, color=TEXT_GREY)
        sp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph().paragraph_format.space_after = Pt(6)


# ── Data loaders ─────────────────────────────────────────────────────────────

def load(filename):
    try:
        with open(os.path.join(DATA_DIR, filename)) as f:
            return json.load(f)
    except Exception as e:
        print(f"  Warning: {filename}: {e}", file=sys.stderr)
        return {}


# ── Section builders ─────────────────────────────────────────────────────────

def build_cover(doc, quarter, updated):
    # Large title block using a borderless table for layout control
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    set_cell_bg(cell, DARK_BLUE)
    cell.width = Inches(6.5)

    for _ in range(2):
        cell.add_paragraph()

    tp = cell.add_paragraph()
    tr = tp.add_run("SUBDIVISION ANALYSIS")
    set_font(tr, bold=True, size=28, color=WHITE)
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sp = cell.add_paragraph()
    sr = sp.add_run("RealtyRates.com Free Survey Data")
    set_font(sr, size=14, color=RGBColor(0xAA, 0xC4, 0xEA))
    sp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for _ in range(1):
        cell.add_paragraph()

    mp = cell.add_paragraph()
    mr = mp.add_run(f"Survey Period: {quarter}  ·  Updated: {updated}")
    set_font(mr, size=10, color=RGBColor(0x88, 0xAA, 0xCC))
    mp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for _ in range(2):
        cell.add_paragraph()

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    src = doc.add_paragraph()
    src.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr2 = src.add_run("Source: realtyrates.com  ·  For informational purposes only")
    set_font(sr2, italic=True, size=8.5, color=RGBColor(0x88, 0x88, 0x88))

    doc.add_page_break()


def build_summary(doc, commercial, developer, indices):
    heading(doc, "Executive Summary", level=1)
    hr(doc)

    sub_cat = next(
        (c for c in developer.get("categories", []) if "Subdivision" in c.get("name", "")), {}
    )
    def get_metric(cat, keyword):
        return next(
            (m for m in cat.get("metrics", []) if keyword.lower() in m.get("name", "").lower()),
            {}
        )

    land  = get_metric(sub_cat, "Land Acquisition")
    ltv   = get_metric(sub_cat, "Max LTV - Land")
    absrp = get_metric(sub_cat, "Absorption")
    eq    = get_metric(sub_cat, "Equity")
    tsy   = next((i for i in indices.get("indices", []) if "10-Year" in i.get("name", "")), {})

    stats = [
        ("LAND LOAN RATE",    fmt_range(land.get("low"), land.get("high")),   f"Avg {fmt_pct(land.get('avg'))}"),
        ("MAX LTV (LAND)",    fmt_range(ltv.get("low"),  ltv.get("high")),    f"Avg {fmt_pct(ltv.get('avg'))}"),
        ("ABSORPTION RATE",   f"{absrp.get('avg', '—')} lots/mo",             f"Range {absrp.get('low','—')}–{absrp.get('high','—')}"),
        ("10-YR TREASURY",    fmt_pct(tsy.get("value")),                      "Benchmark rate"),
    ]
    stat_table(doc, stats)

    body(doc,
         "This report summarises free survey data published by RealtyRates.com covering "
         "commercial mortgage cap rates, developer financing terms for subdivisions and "
         "planned unit developments (PUDs), and key financial benchmark indices. "
         "Data is refreshed quarterly.",
         after=6)

    quarter = commercial.get("quarter") or developer.get("quarter") or "N/A"
    pts = commercial.get("property_types", [])
    if pts:
        cap_avgs = [p["cap_rate_avg"] for p in pts if p.get("cap_rate_avg")]
        rate_avgs = [p["interest_rate_avg"] for p in pts if p.get("interest_rate_avg")]

        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(f"Survey period: ")
        set_font(r, bold=True, size=10, color=TEXT_GREY)
        r2 = p.add_run(quarter)
        set_font(r2, size=10, color=TEXT_GREY)

        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run("Commercial cap rate range: ")
        set_font(r, bold=True, size=10, color=TEXT_GREY)
        r2 = p.add_run(f"{min(cap_avgs):.2f}% – {max(cap_avgs):.2f}% across {len(pts)} property types")
        set_font(r2, size=10, color=TEXT_GREY)

        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run("Commercial interest rate range: ")
        set_font(r, bold=True, size=10, color=TEXT_GREY)
        r2 = p.add_run(f"{min(rate_avgs):.2f}% – {max(rate_avgs):.2f}% (avg, by property type)")
        set_font(r2, size=10, color=TEXT_GREY)

        if land.get("avg") and tsy.get("value"):
            spread = round(land["avg"] - tsy["value"], 2)
            p = doc.add_paragraph(style="List Bullet")
            r = p.add_run("Subdivision land loan spread over 10-Yr Treasury: ")
            set_font(r, bold=True, size=10, color=TEXT_GREY)
            r2 = p.add_run(f"{spread:.2f}% (development risk premium)")
            set_font(r2, size=10, color=TEXT_GREY)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def build_cap_rates(doc, commercial):
    doc.add_page_break()
    heading(doc, "Commercial Cap Rates & Mortgage Rates", level=1)
    hr(doc)
    caption(doc, f"Source: RealtyRates.com Commercial Mortgage Rate Survey  ·  {commercial.get('quarter', '')}")

    pts = commercial.get("property_types", [])
    if not pts:
        body(doc, "No data available.")
        return

    cols = ["Property Type", "Cap Rate Range", "Cap Avg", "Interest Rate Range", "Rate Avg", "Max LTV", "Amort."]
    tbl = doc.add_table(rows=1 + len(pts), cols=len(cols))
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

    # set column widths
    widths = [Inches(2.1), Inches(1.1), Inches(0.7), Inches(1.1), Inches(0.7), Inches(0.6), Inches(0.6)]
    for i, row in enumerate(tbl.rows):
        for j, cell in enumerate(row.cells):
            cell.width = widths[j]

    # header row
    hrow = tbl.rows[0]
    shade_row(hrow, MID_BLUE)
    for j, col in enumerate(cols):
        cell = hrow.cells[j]
        p = cell.paragraphs[0]
        run = p.add_run(col)
        set_font(run, bold=True, size=8.5, color=WHITE)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j > 0 else WD_ALIGN_PARAGRAPH.LEFT

    # data rows
    for i, pt in enumerate(pts):
        row = tbl.rows[i + 1]
        if i % 2 == 1:
            shade_row(row, LIGHT_GREY)

        values = [
            pt["name"],
            fmt_range(pt.get("cap_rate_low"), pt.get("cap_rate_high")),
            fmt_pct(pt.get("cap_rate_avg")),
            fmt_range(pt.get("interest_rate_low"), pt.get("interest_rate_high")),
            fmt_pct(pt.get("interest_rate_avg")),
            fmt_pct(pt.get("ltv_max"), 0) if pt.get("ltv_max") else "—",
            f"{int(pt['amortization'])} yr" if pt.get("amortization") else "—",
        ]
        colors = [TEXT_GREY, TEXT_GREY, ACCENT_GRN, TEXT_GREY, ACCENT_ORG, TEXT_GREY, TEXT_GREY]
        bolds  = [False,     False,     True,        False,     True,        False,     False]

        for j, (val, col_color, bold) in enumerate(zip(values, colors, bolds)):
            cell = row.cells[j]
            p = cell.paragraphs[0]
            run = p.add_run(val)
            set_font(run, bold=bold, size=8.5, color=col_color)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j > 0 else WD_ALIGN_PARAGRAPH.LEFT

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Mini chart placeholder note
    body(doc,
         "[ Chart — Average Cap Rate & Interest Rate by Property Type ]",
         italic=True, size=8.5, color=RGBColor(0x99, 0x99, 0x99), after=2)
    body(doc,
         "Note: Cap Rate Avg (blue) and Interest Rate Avg (orange) are plotted per property type "
         "in the interactive HTML report. Cap rates reflect market-derived yields on stabilised "
         "income-producing properties; interest rates reflect typical permanent loan pricing.",
         size=8.5, color=RGBColor(0x66, 0x66, 0x66), after=6)


def build_developer_survey(doc, developer):
    doc.add_page_break()
    heading(doc, "Developer Survey — Subdivision & Development Financing", level=1)
    hr(doc)
    caption(doc, f"Source: RealtyRates.com Developer Survey  ·  {developer.get('quarter', '')}")

    body(doc,
         "The Developer Survey captures financing terms for residential and commercial development "
         "projects. Unlike stabilised commercial properties, development loans carry higher rates "
         "and tighter LTV constraints to compensate lenders for construction and absorption risk.",
         size=9.5, after=10)

    for cat in developer.get("categories", []):
        heading(doc, cat["name"], level=2, color=MID_BLUE, size=12, after=3)
        if cat.get("description"):
            body(doc, cat["description"], italic=True, size=9, color=TEXT_GREY, after=4)

        metrics = cat.get("metrics", [])
        if not metrics:
            continue

        tbl = doc.add_table(rows=1 + len(metrics), cols=4)
        tbl.style = "Table Grid"
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

        widths = [Inches(2.5), Inches(1.2), Inches(1.2), Inches(1.1)]
        for row in tbl.rows:
            for j, cell in enumerate(row.cells):
                cell.width = widths[j]

        # header
        hrow = tbl.rows[0]
        shade_row(hrow, LIGHT_BLUE)
        for j, hdr in enumerate(["Metric", "Low", "High", "Average"]):
            cell = hrow.cells[j]
            p = cell.paragraphs[0]
            run = p.add_run(hdr)
            set_font(run, bold=True, size=9, color=MID_BLUE)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j > 0 else WD_ALIGN_PARAGRAPH.LEFT

        # data rows
        for i, m in enumerate(metrics):
            row = tbl.rows[i + 1]
            if i % 2 == 1:
                shade_row(row, LIGHT_GREY)
            unit = m.get("unit", "%")
            unit_sym = "%" if unit == "%" else f" {unit}"
            vals = [
                m["name"],
                f"{m.get('low', '—')}{unit_sym}",
                f"{m.get('high', '—')}{unit_sym}",
                f"{m.get('avg', '—')}{unit_sym}",
            ]
            for j, val in enumerate(vals):
                cell = row.cells[j]
                p = cell.paragraphs[0]
                run = p.add_run(val)
                bold = j == 3
                color = ACCENT_GRN if j == 3 else TEXT_GREY
                set_font(run, bold=bold, size=9, color=color)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j > 0 else WD_ALIGN_PARAGRAPH.LEFT

        doc.add_paragraph().paragraph_format.space_after = Pt(8)


def build_indices(doc, indices):
    doc.add_page_break()
    heading(doc, "Financial Indices & Benchmark Rates", level=1)
    hr(doc)
    caption(doc, "Source: RealtyRates.com Financial Indices")

    body(doc,
         "Benchmark rates establish the pricing floor for all real estate lending. "
         "The 10-Year Treasury is the primary reference for long-term fixed-rate commercial "
         "mortgages; Prime Rate and SOFR anchor construction credit lines and floating-rate loans.",
         size=9.5, after=10)

    idx_list = indices.get("indices", [])
    if not idx_list:
        body(doc, "No data available.")
        return

    tbl = doc.add_table(rows=1 + len(idx_list), cols=3)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

    widths = [Inches(2.8), Inches(1.2), Inches(2.5)]
    for row in tbl.rows:
        for j, cell in enumerate(row.cells):
            cell.width = widths[j]

    hrow = tbl.rows[0]
    shade_row(hrow, MID_BLUE)
    for j, hdr in enumerate(["Index", "Current Rate", "Category / Notes"]):
        cell = hrow.cells[j]
        p = cell.paragraphs[0]
        run = p.add_run(hdr)
        set_font(run, bold=True, size=9, color=WHITE)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j == 1 else WD_ALIGN_PARAGRAPH.LEFT

    for i, idx in enumerate(idx_list):
        row = tbl.rows[i + 1]
        if i % 2 == 1:
            shade_row(row, LIGHT_GREY)
        vals = [
            idx["name"],
            fmt_pct(idx.get("value")),
            f"{idx.get('category', '')}  {idx.get('note', '')}".strip(),
        ]
        for j, val in enumerate(vals):
            cell = row.cells[j]
            p = cell.paragraphs[0]
            run = p.add_run(val)
            bold = j == 1
            color = MID_BLUE if j == 1 else TEXT_GREY
            set_font(run, bold=bold, size=9, color=color)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j == 1 else WD_ALIGN_PARAGRAPH.LEFT

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    body(doc,
         "[ Chart — Financial Indices (horizontal bar) available in the interactive HTML report ]",
         italic=True, size=8.5, color=RGBColor(0x99, 0x99, 0x99), after=6)


def build_key_findings(doc, commercial, developer, indices):
    doc.add_page_break()
    heading(doc, "Subdivision Analysis — Key Findings", level=1)
    hr(doc)

    sub_cat = next(
        (c for c in developer.get("categories", []) if "Subdivision" in c.get("name", "")), {}
    )
    def gm(keyword):
        return next(
            (m for m in sub_cat.get("metrics", []) if keyword.lower() in m.get("name", "").lower()), {}
        )

    land  = gm("Land Acquisition")
    const = gm("Construction")
    ltv_l = gm("Max LTV - Land")
    ltv_c = gm("Max LTV - Construction")
    absrp = gm("Absorption")
    eq    = gm("Equity")
    tsy   = next((i for i in indices.get("indices", []) if "10-Year" in i.get("name", "")), {})
    prime = next((i for i in indices.get("indices", []) if "Prime" in i.get("name", "")), {})

    # Comparison: subdivision land rate vs stabilised commercial avg
    pts = commercial.get("property_types", [])
    comm_rate_avg = (
        round(sum(p["interest_rate_avg"] for p in pts if p.get("interest_rate_avg")) / len(pts), 2)
        if pts else None
    )

    findings = []

    if land.get("avg") and tsy.get("value"):
        spread = round(land["avg"] - tsy["value"], 2)
        findings.append((
            "Development Risk Premium",
            f"Subdivision land loans average {fmt_pct(land.get('avg'))}, representing a "
            f"{spread:.2f}% spread over the 10-Year Treasury ({fmt_pct(tsy.get('value'))}). "
            f"This premium reflects lender compensation for entitlement, absorption, and "
            f"market timing risk inherent in raw land."
        ))

    if land.get("avg") and comm_rate_avg:
        gap = round(land["avg"] - comm_rate_avg, 2)
        findings.append((
            "Subdivision vs. Stabilised Commercial",
            f"Land loan rates ({fmt_pct(land.get('avg'))} avg) run {gap:.2f}% above the average "
            f"permanent loan rate for stabilised commercial properties ({fmt_pct(comm_rate_avg)}). "
            f"Construction loan rates ({fmt_pct(const.get('avg'))}) carry a similar premium."
        ))

    if ltv_l.get("avg"):
        findings.append((
            "Conservative LTV Constraints",
            f"Maximum LTV on land acquisition averages {fmt_pct(ltv_l.get('avg'))} "
            f"(range {fmt_range(ltv_l.get('low'), ltv_l.get('high'))}), compared to 65–75% "
            f"typical of stabilised commercial assets. Construction LTV averages "
            f"{fmt_pct(ltv_c.get('avg'))}. Developers must bring meaningful equity to the table."
        ))

    if absrp.get("avg"):
        findings.append((
            "Absorption Rate is the Primary Feasibility Lever",
            f"Lenders typically underwrite {absrp.get('avg')} lots/month (range "
            f"{absrp.get('low')}–{absrp.get('high')} lots/month). Faster absorption reduces "
            f"carry cost and shortens the loan term, directly improving project IRR. "
            f"Market demand, pricing strategy, and amenity package all drive this metric."
        ))

    if eq.get("avg"):
        findings.append((
            "Equity Requirement",
            f"Lenders require {fmt_pct(eq.get('avg'))} equity on average "
            f"(range {fmt_range(eq.get('low'), eq.get('high'))}). "
            f"With Prime Rate at {fmt_pct(prime.get('value'))}, carry costs are elevated — "
            f"making pre-sales and phased releases critical to managing cash flow."
        ))

    findings.append((
        "Rate Environment Context",
        f"With the Federal Funds Rate at ~4.375% and the 10-Year Treasury at "
        f"{fmt_pct(tsy.get('value'))}, land loan rates remain historically elevated. "
        f"Projects underwritten at today's rates should stress-test absorption assumptions "
        f"and plan for potential refinancing once rates moderate."
    ))

    for i, (title, text) in enumerate(findings):
        heading(doc, f"{i+1}.  {title}", level=2, color=DARK_BLUE, size=11, after=2)
        body(doc, text, size=9.5, color=TEXT_GREY, after=8)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def build_methodology(doc):
    doc.add_page_break()
    heading(doc, "Methodology & Disclaimers", level=1)
    hr(doc)

    paras = [
        ("Data Source",
         "All data is sourced from the free survey publications available at "
         "realtyrates.com. RealtyRates.com aggregates quarterly responses from over "
         "300 appraisers, brokers, and direct lenders across the United States."),
        ("Survey Coverage",
         "Three free surveys are captured: (1) Commercial Mortgage Rate Survey — cap "
         "rates and loan terms for 15+ income-producing property types; "
         "(2) Financial Indices — key benchmark rates; "
         "(3) Developer Survey — financing terms for subdivisions, PUDs, condominiums, "
         "and business/industrial parks."),
        ("Update Frequency",
         "RealtyRates.com publishes survey data quarterly. This report is refreshed "
         "automatically each week via a GitHub Actions workflow; the data version reflects "
         "the most recently published quarter."),
        ("Limitations",
         "Survey data represents typical market ranges and averages — individual "
         "transactions may fall outside these ranges based on sponsor quality, market, "
         "property condition, and lender appetite. This report is for informational "
         "purposes only and does not constitute investment or lending advice."),
    ]
    for title, text in paras:
        p = doc.add_paragraph()
        r1 = p.add_run(f"{title}: ")
        set_font(r1, bold=True, size=9.5, color=DARK_BLUE)
        r2 = p.add_run(text)
        set_font(r2, size=9.5, color=TEXT_GREY)
        p.paragraph_format.space_after = Pt(6)


# ── Page setup & styles ──────────────────────────────────────────────────────

def setup_document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width  = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin    = Inches(0.9)
    sec.bottom_margin = Inches(0.9)
    sec.left_margin   = Inches(1.0)
    sec.right_margin  = Inches(1.0)

    # Default paragraph style
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)
    style.font.color.rgb = TEXT_GREY

    # Heading styles
    for lvl, sz in [(1, 15), (2, 12), (3, 11)]:
        h = doc.styles[f"Heading {lvl}"]
        h.font.name = "Calibri"
        h.font.size = Pt(sz)
        h.font.bold = True
        h.font.color.rgb = DARK_BLUE
        h.paragraph_format.space_before = Pt(12 if lvl == 1 else 8)
        h.paragraph_format.space_after  = Pt(4)

    return doc


# ── Main ─────────────────────────────────────────────────────────────────────

def main() -> int:
    print("Loading data…")
    commercial = load("commercial_rates.json")
    indices    = load("indices.json")
    developer  = load("developer_survey.json")
    metadata   = load("metadata.json")

    quarter = commercial.get("quarter") or developer.get("quarter") or "Latest Quarter"
    last_updated = metadata.get("last_updated", "")
    try:
        dt = datetime.fromisoformat(last_updated.replace("Z", "+00:00"))
        updated_fmt = dt.strftime("%B %d, %Y")
    except Exception:
        updated_fmt = last_updated or "N/A"

    print("Building Word document…")
    doc = setup_document()

    build_cover(doc, quarter, updated_fmt)
    build_summary(doc, commercial, developer, indices)
    build_cap_rates(doc, commercial)
    build_developer_survey(doc, developer)
    build_indices(doc, indices)
    build_key_findings(doc, commercial, developer, indices)
    build_methodology(doc)

    os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
    doc.save(OUT_PATH)
    print(f"Saved: {OUT_PATH}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
